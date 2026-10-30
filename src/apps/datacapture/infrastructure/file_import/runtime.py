"""Technical boundaries for workbook parsing, preview signing and capture locking."""
import hashlib
import json
from contextlib import contextmanager
from datetime import date, datetime, time
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from django.core import signing
from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.datacapture.infrastructure.repositories.page_capture import DjangoDataCapturePageRepository
from apps.datacapture.models import DataCapturePageEntry, DataCapturePageState
from apps.subject.public import get_subject_capture_eligibility

MAX_BYTES = 5 * 1024 * 1024
SALT = "datacapture.mapped-file-import.v1"


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str, ensure_ascii=False).encode()).hexdigest()


def sign_preview(payload):
    return signing.dumps(payload, salt=SALT, compress=True)


def read_preview(token):
    try:
        return signing.loads(token, salt=SALT, max_age=900)
    except signing.BadSignature as exc:
        raise ValueError(_("The preview has expired or is invalid. Generate a new preview.")) from exc


def read_workbook(content):
    if not content or len(content) > MAX_BYTES:
        raise ValueError(_("Choose an .xlsx file no larger than 5 MB."))
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError(_("The uncompressed Excel file exceeds 50 MB."))
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        result = {}
        total_cells = 0
        try:
            for sheet in workbook:
                if not sheet.max_row or not sheet.max_column or sheet.max_row > 20000 or sheet.max_column > 200:
                    raise ValueError(_("The sheet exceeds the limit of 20,000 rows or 200 columns."))
                total_cells += sheet.max_row * sheet.max_column
                if total_cells > 500000:
                    raise ValueError(_("The workbook exceeds 500,000 cells. Split the results into smaller files."))
                rows = []
                for row in sheet.iter_rows():
                    rows.append([{"value": cell.value.isoformat() if isinstance(cell.value, (date, datetime, time)) else cell.value,
                                  "formula": cell.data_type == "f", "excel_error": cell.data_type == "e"} for cell in row])
                result[sheet.title] = rows
            return result
        finally:
            workbook.close()
    except (BadZipFile, KeyError, OSError, TypeError, ParseError) as exc:
        raise ValueError(_("The .xlsx file could not be read.")) from exc


def match_rows(workbook, config):
    preview = config["preview"]
    rows = workbook.get(preview["sheet"], [])
    header_index = preview["header_row"] - 1
    if len(rows) <= header_index:
        return None
    headers = rows[header_index]
    expected = {item["column"]: item["header"] for item in preview["columns"]}
    actual = {}
    for index, cell in enumerate(headers):
        letter = get_column_letter(index + 1)
        if cell["formula"]:
            return None
        value = cell["value"]
        if value not in (None, "") or letter in expected:
            actual[letter] = str(value or f"Cột {letter} chưa có tiêu đề").strip()[:500]
    if actual != expected:
        return None
    return [{"row_number": index + 1, "cells": {get_column_letter(i + 1): cell for i, cell in enumerate(row)}}
            for index, row in enumerate(rows) if index > header_index
            and any(cell["value"] is not None and str(cell["value"]).strip() for cell in row)]


class FileImportCaptureRepository:
    def __init__(self, pages=None):
        self.pages = pages or DjangoDataCapturePageRepository()

    @contextmanager
    def atomic_scope(self, subject_id, visit_id):
        with transaction.atomic():
            eligibility = get_subject_capture_eligibility(subject_id=subject_id, event_instance_id=visit_id, for_update=True)
            if not eligibility.allowed:
                raise ValueError(_("Data entry is not allowed for this visit (%(reason)s).") % {"reason": eligibility.reason})
            list(DataCapturePageState.objects.select_for_update().filter(subject_id=subject_id, visit_id=visit_id).values_list("id", flat=True))
            list(DataCapturePageEntry.objects.select_for_update().filter(subject_id=subject_id, visit_id=visit_id).values_list("id", flat=True))
            yield

    def snapshot(self, subject_id, visit_id, form_id, binding_id):
        scope = dict(subject_id=subject_id, visit_id=visit_id, crf_template_id=form_id, event_form_binding_id=binding_id)
        page = self.pages.get_page_state_by_scope(**scope)
        entry = self.pages.get_current_entry(**scope)
        raw = entry.data if entry else (page.final_data if page else "{}")
        data = self.pages.flatten_form_data_json_for_read(data=raw, crf_template_id=form_id)
        return {"page_id": page.id if page else None, "page_status": page.status if page else None,
                "entry_id": entry.id if entry else None, "entry_status": entry.status if entry else None,
                "updated_at": str(entry.updated_at) if entry else None, "data": data,
                "version": page.data_version if page else None, "raw_data": json.loads(raw or "{}")}
