import re
import unicodedata
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.datacapture.application.services.laboratory_import_mapping import LAB_HEADERS

MAX_UPLOAD_BYTES = 2 * 1024 * 1024


def normalize_header(value):
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split()).casefold()


def parse_laboratory_workbook(source):
    """Read one result row without inferring patient, visit, units or clinical meaning."""
    source.seek(0, 2)
    if source.tell() > MAX_UPLOAD_BYTES:
        raise ValueError("File Excel vượt quá 2 MB.")
    source.seek(0)
    try:
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("Nội dung Excel sau giải nén quá lớn.")
        source.seek(0)
        workbook = load_workbook(source, read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("File không phải Excel .xlsx hợp lệ.") from exc
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("File phải có đúng một sheet kết quả.")
        sheet = workbook.worksheets[0]
        if sheet.max_row > 100 or sheet.max_column > 100:
            raise ValueError("File vượt giới hạn 100 dòng hoặc 100 cột.")
        rows = list(sheet.iter_rows())
        if len(rows) < 3:
            raise ValueError("Thiếu dòng tiêu đề thứ 2 hoặc dòng kết quả.")
        headers = rows[1]
        data_rows = [row for row in rows[2:] if any(c.value not in (None, "") for c in row)]
        if len(data_rows) != 1:
            raise ValueError("File phải có đúng một dòng kết quả; không tự chọn giữa nhiều người/mẫu.")
        mapping = {normalize_header(header): key for header, key in LAB_HEADERS.items()}
        seen = set()
        results = []
        for index, (header, cell) in enumerate(zip(headers, data_rows[0]), start=1):
            normalized = normalize_header(header.value)
            if normalized and normalized in seen:
                raise ValueError(f"Tiêu đề bị trùng: {header.value}.")
            seen.add(normalized)
            if cell.value in (None, ""):
                continue
            if cell.data_type == "f" or header.data_type == "f":
                raise ValueError(f"Ô {cell.coordinate} chứa công thức; cần giá trị kết quả gốc.")
            key = mapping.get(normalized)
            raw = str(cell.value).strip()
            item = {
                "cell": f"{get_column_letter(index)}{cell.row}",
                "header": str(header.value or "(Không có tiêu đề)"),
                "raw": raw,
                "field_key": key,
                "value": None,
                "warning": "",
            }
            if key is None:
                item["warning"] = "Chưa có trường CRF tương ứng; không nhập."
            else:
                item["value"] = convert_result(key, raw)
            results.append(item)
        if not any(item["field_key"] for item in results):
            raise ValueError("Không nhận diện được chỉ số xét nghiệm trong file.")
        return results
    finally:
        workbook.close()


def convert_result(key, raw):
    if key.endswith("_RESULT") or (key.startswith("LAB_URINE_") and key not in {"LAB_URINE_PH", "LAB_URINE_SG"}):
        qualitative = normalize_header(raw)
        if key.endswith("_RESULT"):
            match = re.fullmatch(r"(?:\d+(?:[.,]\d+)?\s*\()?((?:âm|dương) tính)\)?", qualitative)
            if match:
                return {"âm tính": "0", "dương tính": "1"}[match[1]]
        elif qualitative in {"âm tính", "dương tính", "vết"}:
            return {"âm tính": "1", "dương tính": "2", "vết": "3"}[qualitative]
        raise ValueError(f"{key}: không nhận diện kết quả định tính '{raw}'.")
    if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", raw):
        raise ValueError(f"{key}: '{raw}' không phải số hợp lệ; không tự bỏ ký hiệu hoặc đơn vị.")
    try:
        value = Decimal(raw.replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"{key}: số không hợp lệ.") from exc
    return format(value, "f")
