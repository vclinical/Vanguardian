from datetime import date, datetime
from io import BytesIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 5 * 1024 * 1024


def read_import_workbook(content, *, sheet_name="", header_row=2):
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("Chọn file Excel không quá 5 MB.")
    if not isinstance(header_row, int) or not 1 <= header_row <= 100:
        raise ValueError("Dòng tiêu đề phải từ 1 đến 100.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Nội dung Excel sau giải nén vượt quá 50 MB.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, KeyError, ParseError) as exc:
        raise ValueError("Không đọc được file .xlsx. Hãy kiểm tra file Excel gốc.") from exc
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError("Sheet được chọn không tồn tại.")
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        if not sheet.max_row or not sheet.max_column or sheet.max_row > 20000 or sheet.max_column > 200:
            raise ValueError("Sheet phải có tối đa 20.000 dòng và 200 cột, với vùng dữ liệu xác định.")
        if header_row >= sheet.max_row:
            raise ValueError("Không có dữ liệu sau dòng tiêu đề đã chọn.")
        headers = next(sheet.iter_rows(min_row=header_row, max_row=header_row))
        if any(cell.data_type == "f" for cell in headers):
            raise ValueError("Tên cột không được là công thức Excel.")
        samples = [[] for _ in headers]
        populated = set()
        row_count = 0
        for row in sheet.iter_rows(min_row=header_row + 1):
            if not any(cell.value not in (None, "") for cell in row):
                continue
            row_count += 1
            for index, cell in enumerate(row):
                value = cell.value
                if value in (None, ""):
                    continue
                populated.add(index)
                if len(samples[index]) < 5:
                    value = value.isoformat() if isinstance(value, (date, datetime)) else str(value)
                    samples[index].append({"cell": cell.coordinate, "value": value[:500], "formula": cell.data_type == "f"})
        columns = []
        for index, cell in enumerate(headers):
            if cell.value in (None, "") and index not in populated:
                continue
            letter = get_column_letter(index + 1)
            columns.append({
                "column": letter,
                "header": str(cell.value or f"Cột {letter} chưa có tiêu đề").strip()[:500],
                "samples": samples[index],
            })
        if not columns or not row_count:
            raise ValueError("Không tìm thấy cột hoặc dòng dữ liệu. Hãy kiểm tra sheet và dòng tiêu đề.")
        return {"sheet": sheet.title, "sheets": workbook.sheetnames, "header_row": header_row,
                "row_count": row_count, "columns": columns}
    except (ParseError, BadZipFile, KeyError) as exc:
        raise ValueError("Nội dung sheet không hợp lệ.") from exc
    finally:
        workbook.close()
